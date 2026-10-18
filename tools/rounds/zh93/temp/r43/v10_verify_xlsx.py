# -*- coding: utf-8 -*-
r"""v10_verify_xlsx.py — 把 字库报告.xlsx 读回来复核 (验收③)

· 页数与每页行数;
· 概要页的关键数字 vs **dump_font 自己产出的** summary.json (必须一致);
· 空位清单页必须正好 45 行 = summary.json 的 vacuum_usable;
· 手柄图标码页必须有 0xB1/0xB5/0xE7/0xF1 四个字节, 且 0xE7/0xF1 标"不是图标"。
"""
import io
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                    '..', '..', '..', '..', '..'))
REL = os.path.join(REPO, 'work', 'builds', 'GR_ZH93', 'release')
XLSX = os.path.join(REL, '字库报告.xlsx')
SUMJ = os.path.join(REPO, 'work', 'builds', 'GR_ZH93', 'font', 'GR_ZH93_font_summary.json')

from openpyxl import load_workbook                                     # noqa: E402

ok = True


def chk(cond, msg):
    global ok
    ok = ok and bool(cond)
    print('%s %s' % ('[ok]' if cond else '[!!]', msg))


summ = json.load(io.open(SUMJ, encoding='utf-8'))
wb = load_workbook(XLSX, data_only=True)
print('# 文件: %s (%d B)' % (os.path.relpath(XLSX, REPO).replace('\\', '/'),
                             os.path.getsize(XLSX)))
print('# 页 (%d): %s' % (len(wb.sheetnames), ' / '.join(wb.sheetnames)))
chk(wb.sheetnames == ['索引', '概要', '字符分类', '空位清单', '手柄图标码', '口径与限制'],
    '页序与页名符合设计 (索引/概要/字符分类/空位清单/手柄图标码/口径与限制)')
for nm in wb.sheetnames:
    ws = wb[nm]
    print('  %-8s %3d 行 × %2d 列  (A1=%r)' % (nm, ws.max_row, ws.max_column,
                                               str(ws['A1'].value)[:28]))

# ── 概要页: 关键数字交叉复核 ──
ws = wb['概要']
kv = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r and r[0]:
        kv[str(r[0])] = r[1]
c = summ['capacity']
p = c['pair']
want = {
    '【汉字·对偶码】总容量': p['capacity'],
    '【汉字·对偶码】已用': p['used'],
    '【汉字·对偶码】空': p['free'],
    '★ 可直接用（新字马上能加）': p['vacuum_usable'],
    '需回收（格子上有旧字）': p['need_recycle'],
    '单字节码 已用（总）': c['single']['used'],
    '单字节码 ★真正可用': c['single']['really_usable'],
    '字符总数（charset 口径）': c['charset']['unique_chars'],
    '字符总数（dump 口径）': c['charset']['dump_total'],
}
print('\n# 概要页 vs font_summary.json (dump_font 现算):')
for k, v in want.items():
    got = kv.get(k)
    chk(got == v, '%-28s 报告=%s  dump_font=%s' % (k, got, v))
chk(str(kv.get('图集占用')) == '%.2f%%' % c['atlas']['ink_ratio'],
    '图集占用 %s == %.2f%%' % (kv.get('图集占用'), c['atlas']['ink_ratio']))
chk(str(kv.get('对偶格占用')) == '%.2f%%' % c['atlas']['pair_cell_ratio'],
    '对偶格占用 %s == %.2f%%' % (kv.get('对偶格占用'), c['atlas']['pair_cell_ratio']))

# ── 空位清单页 ──
ws = wb['空位清单']
rows = list(ws.iter_rows(min_row=2, values_only=True))
chk(len(rows) == p['vacuum_usable'], '空位清单 %d 行 == dump_font 的 vacuum_usable %d'
    % (len(rows), p['vacuum_usable']))
chk([r[0] for r in rows] == p['vacuum_slots'],
    '空位清单的 slot 序列与 dump_font 的 vacuum_slots 逐个相同')
bad = [r for r in rows if not (0 <= r[0] < p['capacity'])]
chk(not bad, '所有 slot 都在 0..%d 内' % (p['capacity'] - 1))
chk(all(isinstance(r[3], int) and isinstance(r[4], int) for r in rows),
    'cell_x/cell_y 是数字 (Excel 里可直接排序/筛选)')
n_strict = len([r for r in rows if r[9] == '是'])
chk(n_strict == p['strict_window_free'],
    '严格口径干净的 %d 个 == dump_font 的 strict_window_free %d'
    % (n_strict, p['strict_window_free']))
print('  前 3 行: %s' % (rows[:3],))

# ── 字符分类页 ──
ws = wb['字符分类']
kinds = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r and r[0] in ('ascii', 'single', 'latin1', 'pair', 'icon'):
        kinds[r[0]] = r[1]
chk(kinds == {'ascii': 95, 'single': 108, 'latin1': 2, 'pair': 1319, 'icon': 2},
    '字符分类页: %s (latin1=2, icon=2 是本轮硬指标)' % kinds)

# ── 手柄图标码页 ──
ws = wb['手柄图标码']
txt = '\n'.join(str(c) for r in ws.iter_rows(values_only=True) for c in r if c)
for t in ('0xB1', '0xB5', '0xE7', '0xF1', '三角键', '叉键', '不是图标', '历史别名'):
    chk(t in txt, '手柄图标码页含 %r' % t)

print('\n结论: %s' % ('[ok] xlsx 读回复核全部通过 ✓' if ok else '[!!] 有失败项 ✗'))
sys.exit(0 if ok else 1)
