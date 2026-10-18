# -*- coding: utf-8 -*-
r"""text_table_xlsx.py — 把 dump_all_text.py 的 all_text.csv 渲染成多 sheet 的 Excel 总表

用法: python tools/diag/text_table_xlsx.py <textdump_dir> [--out xlsx路径]

产出 sheet:
  0 总览        各来源条数/中文率/差异数
  1 全部文本    全量行 (source/key/en/zh/changed), 冻结首行 + 自动筛选 + 差异高亮
  2..N 分来源   每个来源一张表 (便于按资源排查)
"""
import argparse
import collections
import csv
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print('需要 openpyxl: pip install openpyxl')
    sys.exit(2)

HDR = ['source', 'key', 'en', 'zh', 'changed', 'len_en', 'len_zh', 'offset']
WIDTH = [16, 26, 58, 58, 9, 8, 8, 12]
SAFE = {'\\': '＼', '/': '／', '?': '？', '*': '＊', '[': '［', ']': '］', ':': '：'}


def sheet_name(s, used):
    t = ''.join(SAFE.get(ch, ch) for ch in s)[:31] or 'sheet'
    base, k = t, 2
    while t in used:
        t = '%s_%d' % (base[:28], k)
        k += 1
    used.add(t)
    return t


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='3A4A6B')
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(c)].width = WIDTH[c - 1]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(ncol), ws.max_row)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('dumpdir')
    ap.add_argument('--out', default=None)
    a = ap.parse_args()
    src = os.path.join(a.dumpdir, 'all_text.csv')
    assert os.path.isfile(src), src
    rows = list(csv.DictReader(open(src, encoding='utf-8-sig')))
    out = a.out or os.path.join(a.dumpdir, 'all_text.xlsx')

    wb = Workbook()
    # --- 总览 ---
    ws = wb.active
    ws.title = '总览'
    by = collections.OrderedDict()
    for r in rows:
        s = by.setdefault(r['source'], [0, 0, 0])
        s[0] += 1
        if any('\u4e00' <= ch <= '\u9fff' for ch in r['zh']):
            s[1] += 1
        if r['changed']:
            s[2] += 1
    ws.append(['来源', '条数', '含中文', '中文率', '与本盘差异'])
    tot = [0, 0, 0]
    for s, (n, m, c) in by.items():
        ws.append([s, n, m, round(100.0 * m / n, 1) if n else 0, c])
        tot[0] += n; tot[1] += m; tot[2] += c
    ws.append(['合计', tot[0], tot[1], round(100.0 * tot[1] / tot[0], 1) if tot[0] else 0, tot[2]])
    for c in range(1, 6):
        ws.cell(row=1, column=c).font = Font(bold=True, color='FFFFFF')
        ws.cell(row=1, column=c).fill = PatternFill('solid', fgColor='3A4A6B')
        ws.column_dimensions[get_column_letter(c)].width = [22, 10, 10, 10, 12][c - 1]
    ws.freeze_panes = 'A2'
    for c in range(1, 6):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)

    # --- 全部文本 ---
    ws = wb.create_sheet('全部文本')
    ws.append(HDR)
    hl = PatternFill('solid', fgColor='E8F4E8')
    for r in rows:
        ws.append([r['source'], r['key'], r['en'], r['zh'], r['changed'],
                   int(r['len_en'] or 0), int(r['len_zh'] or 0), r['offset']])
        if r['changed']:
            for c in range(1, len(HDR) + 1):
                ws.cell(row=ws.max_row, column=c).fill = hl
    style_header(ws, len(HDR))

    # --- 分来源 ---
    used = {'总览', '全部文本'}
    groups = collections.OrderedDict()
    for r in rows:
        groups.setdefault(r['source'], []).append(r)
    for s, rs in groups.items():
        ws = wb.create_sheet(sheet_name(s, used))
        ws.append(HDR)
        for r in rs:
            ws.append([r['source'], r['key'], r['en'], r['zh'], r['changed'],
                       int(r['len_en'] or 0), int(r['len_zh'] or 0), r['offset']])
            if r['changed']:
                for c in range(1, len(HDR) + 1):
                    ws.cell(row=ws.max_row, column=c).fill = hl
        style_header(ws, len(HDR))

    wb.save(out)
    print('xlsx 写出: %s (%d 行, %d 个 sheet)'
          % (out, len(rows), len(wb.sheetnames)))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    main()
